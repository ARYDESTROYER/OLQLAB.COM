#!/usr/bin/env python3
import json
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl


def normalize_section_title(raw: str) -> str:
    if raw.lower().startswith("section 1"):
        return "Cognitive Orientation"
    if raw.lower().startswith("section 2"):
        return "Personality Orientation"
    if raw.lower().startswith("section 3"):
        return "Response Orientation"
    if raw.lower().startswith("the probablity dilemma"):
        return "The Probability Dilemma"
    if raw.lower().startswith("word association"):
        return "Word Association"
    if raw.lower().startswith("tale spin"):
        return "Tale Spin"
    return raw.strip()


def build_payload(xlsx_path: Path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    sections: "OrderedDict[str, dict]" = OrderedDict()
    current_section = "Cognitive Orientation"

    for row in range(1, ws.max_row + 1):
        qno = ws.cell(row, 2).value
        prompt = ws.cell(row, 3).value
        options = [ws.cell(row, col).value for col in range(4, 9)]

        if isinstance(qno, str) and qno.strip():
            value = qno.strip()
            if value.lower().startswith("section") or value.lower() in {
                "the probablity dilemma",
                "word association",
                "tale spin",
            }:
                current_section = normalize_section_title(value)
                if current_section not in sections:
                    sections[current_section] = {
                        "title": current_section,
                        "kind": "SCENARIO",
                        "questions": [],
                    }
                continue

        if not isinstance(qno, int):
            continue

        if current_section not in sections:
            sections[current_section] = {
                "title": current_section,
                "kind": "SCENARIO",
                "questions": [],
            }

        text_prompt = str(prompt or "").strip()
        option_text = [str(opt).strip() for opt in options if isinstance(opt, str) and opt.strip()]

        if current_section == "Word Association":
            question = {
                "code": f"Q{qno}",
                "prompt": text_prompt,
                "questionType": "FREE_TEXT",
                "category": current_section,
                "reverse": False,
                "scaleMin": 1,
                "scaleMax": 5,
            }
        else:
            question = {
                "code": f"Q{qno}",
                "prompt": text_prompt,
                "questionType": "SJT_SINGLE",
                "category": current_section,
                "reverse": False,
                "scaleMin": 1,
                "scaleMax": 5,
                "options": [
                    {
                        "code": chr(65 + idx),
                        "text": option,
                    }
                    for idx, option in enumerate(option_text)
                ],
            }

        sections[current_section]["questions"].append(question)

    return {
        "title": "Wissen Leadership Readiness Test (CPR)",
        "competencies": [],
        "sections": list(sections.values()),
        "policy": {
            "showResultsToEmployee": True,
            "resultReleaseDelayHours": 0,
            "postSubmitMessage": "Assessment completed. You will be notified once your report is available.",
            "leaderCanViewFullReport": True,
            "reportWorkflow": "MANUAL_PDF_UPLOAD",
            "randomizeQuestionOrder": False,
            "submissionAlertAdminIds": [],
        },
    }


def main():
    if len(sys.argv) < 3:
        print(
            "Usage: generate-cpr-manual-upload-payload.py <input.xlsx> <output.json>",
            file=sys.stderr,
        )
        sys.exit(1)

    input_path = Path(sys.argv[1]).expanduser().resolve()
    output_path = Path(sys.argv[2]).expanduser().resolve()

    payload = build_payload(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"Wrote CPR manual upload payload to {output_path}")


if __name__ == "__main__":
    main()
